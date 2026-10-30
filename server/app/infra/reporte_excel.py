"""El reporte en Excel para el administrador (ADR 0002 §127).

────────────────────────────────────────────────────────────────────────────
PARA QUIÉN
────────────────────────────────────────────────────────────────────────────
«Quiero generar un reporte en Excel para la persona que usa el perfil de
administrador; él no sabe usar mucho sistemas, así que el dashboard solo lo uso
yo, pero el reporte se lo daré a él.» Eligió las cuatro cosas —ventas y
efectivo, inventario, cargas y mermas, y la cuenta de cada vendedor—, del
periodo que se elija al bajarlo.

Así que el archivo se lee sin saber nada del sistema: una hoja de **Resumen**
con las cifras grandes en palabras, y una hoja por tema con sus totales al
final. Encabezados en rojo, pesos con signo, fechas con día, y nada de claves
internas ni de estados en inglés.

────────────────────────────────────────────────────────────────────────────
LAS MISMAS CIFRAS QUE EL DASHBOARD
────────────────────────────────────────────────────────────────────────────
Las ventas son las confirmadas, como el tablero; lo eliminado en Operaciones
(§119, §126) no cuenta en ninguna hoja. El efectivo es el de las ventas de
contado en efectivo, que es el que decide el corte. Y como toda venta es en
efectivo (§81, §154), lo vendido no se repite como «en efectivo»: el renglón y
las columnas del efectivo salen solo si en el periodo hubo una transferencia
de un teléfono anterior a la +28, porque entonces ya no son la misma cifra
(§170). El inventario es el de hoy,
valuado a precio de venta como en Inventario (§86): no depende del periodo. Lo
que valía al cerrar cada día del periodo, la bodega y cada camión, y en qué se
fue o de dónde vino lo que cambió, está en «Día por día» (§183, §184), a los
mismos precios.
La cuenta del vendedor es su saldo de hoy, y sus movimientos del periodo.
"""

from __future__ import annotations

import io
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from app.infra.corte_de_la_carga import DESDE_LA_CARGA, TOPE_DEL_CORTE
from app.infra.cuenta_vendedor import ORIGENES
from app.infra.valor_de_inventario import precio_de_venta

ROJO = "E0282E"
GRIS = "F2F2F2"
DINERO = '"$"#,##0.00'
# Sin centavos, donde las columnas son angostas.
PESOS = '"$"#,##0'
ENTERO = "#,##0"
FECHA = "dd/mm/yyyy"

_DIAS = ("lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo")

ESTADO_DE_LA_CARGA = {
    "confirmada": "En el camión",
    "en_ruta": "En el corte",
    "liquidada": "Cortada",
}

TIPO_DE_MOVIMIENTO = {
    "cargo": "Se le cargó",
    "abono": "Pagó",
    "condonacion": "Se le perdonó",
}


@dataclass
class Reporte:
    desde: object
    hasta: object
    descripcion: str
    ventas: list[dict] = field(default_factory=list)
    camiones: list[dict] = field(default_factory=list)
    mayoreo: list[dict] = field(default_factory=list)
    cortes: list[dict] = field(default_factory=list)
    inventario: list[dict] = field(default_factory=list)
    cargas: list[dict] = field(default_factory=list)
    mermas: list[dict] = field(default_factory=list)
    saldos: list[dict] = field(default_factory=list)
    movimientos: list[dict] = field(default_factory=list)
    valor_por_dia: dict = field(default_factory=dict)


# ---------------------------------------------------------------------------
# Los datos
# ---------------------------------------------------------------------------

SQL_VENTAS = """
SELECT v.fecha_operativa AS fecha, u.nombre AS vendedor,
       count(*) AS ventas,
       COALESCE(sum(p.piezas), 0) AS piezas,
       sum(v.total) AS importe,
       COALESCE(sum(v.total) FILTER (WHERE v.tipo = 'contado' AND v.forma_pago = 'efectivo'),
                0) AS efectivo
  FROM ventas v
  JOIN usuarios u ON u.id = v.vendedor_id
  LEFT JOIN LATERAL (
        SELECT sum(cantidad_base) AS piezas FROM venta_partidas WHERE venta_id = v.id
  ) p ON true
 WHERE v.fecha_operativa BETWEEN :desde AND :hasta AND v.estado = 'confirmada'
 GROUP BY v.fecha_operativa, u.nombre
 ORDER BY v.fecha_operativa, u.nombre
"""

# «Quiero que aparezcan las ventas de cada día de cada vendedor: vendió tantos de
# este artículo, tantos de estos», «y que en el mismo día venga la carga que hizo,
# para llevar un control del inventario día a día» (ADR 0002 §177). Y después:
# «quiero poder comparar por vendedor día a día, ver cómo se comporta su
# inventario» (§179). Así que va por camión, que es donde está la mercancía: lo
# que traía cada día, lo que cargó, vendió y mermó, lo demás, y lo que le quedó.

# La carga va el día en que se vendió con ella: el primero con ventas de los que
# cuenta su corte (§176), o su propio día si no vendió. La del sábado queda «para
# el domingo» y se vende el lunes: sale el lunes.
_DIAS_DE_LA_CARGA = f"""
    SELECT c.id, c.almacen_destino_id AS camion,
           COALESCE((SELECT min(v.fecha_operativa) FROM ventas v
                      WHERE v.vendedor_id = c.vendedor_id AND v.estado = 'confirmada'
                        AND v.fecha_operativa >= {DESDE_LA_CARGA}
                        AND v.fecha_operativa < {TOPE_DEL_CORTE}),
                    c.fecha_operativa) AS dia
      FROM cargas c
     WHERE c.estado IN ('confirmada', 'en_ruta', 'liquidada')
"""

SQL_CAMIONES = """
SELECT a.id AS camion, a.nombre, u.nombre AS vendedor
  FROM almacenes a
  LEFT JOIN usuarios u ON u.id = a.responsable_id
 WHERE a.tipo = 'camion'
"""

SQL_CARGADO_AL_CAMION = f"""
WITH dias AS ({_DIAS_DE_LA_CARGA})
SELECT k.camion, k.dia, d.producto_id AS producto, sum(d.cantidad) AS cantidad
  FROM dias k
  JOIN carga_detalle d ON d.carga_id = k.id
 WHERE k.dia BETWEEN :desde AND :hasta
 GROUP BY k.camion, k.dia, d.producto_id
"""

# Las mismas ventas que la hoja de Ventas —confirmadas—, del camión de donde
# salieron.
SQL_VENDIDO_DEL_CAMION = """
SELECT v.almacen_id AS camion, v.fecha_operativa AS dia, vp.producto_id AS producto,
       sum(vp.cantidad_base) AS cantidad
  FROM ventas v
  JOIN venta_partidas vp ON vp.venta_id = v.id
 WHERE v.fecha_operativa BETWEEN :desde AND :hasta AND v.estado = 'confirmada'
 GROUP BY v.almacen_id, v.fecha_operativa, vp.producto_id
"""

SQL_MERMA_DEL_CAMION = """
SELECT m.almacen_id AS camion, m.fecha_operativa AS dia, d.producto_id AS producto,
       sum(d.cantidad_base) AS cantidad
  FROM mermas m
  JOIN merma_detalle d ON d.merma_id = m.id
 WHERE m.fecha_operativa BETWEEN :desde AND :hasta AND m.estado = 'confirmada'
   AND m.tipo IN ('merma', 'cambio')
 GROUP BY m.almacen_id, m.fecha_operativa, d.producto_id
"""

# Todo lo que entró o salió de cada camión desde el primer día, del libro de
# movimientos, en el día de su documento: la venta y la merma en el suyo, la
# carga en el día en que se vendió con ella; lo demás —ajustes, devoluciones,
# lo que regresó a la bodega— el día en que se hizo. Con lo que el camión tiene
# hoy, de aquí sale con cuánto terminó cada día.
SQL_MOVIMIENTO_DEL_CAMION = f"""
WITH dias AS ({_DIAS_DE_LA_CARGA}),
efectos AS (
    SELECT a.id AS camion, m.producto_id AS producto,
           COALESCE(v.fecha_operativa, me.fecha_operativa, k.dia,
                    CAST(m.fecha_servidor AS date)) AS dia,
           CASE WHEN m.almacen_destino_id = a.id THEN m.cantidad
                ELSE -m.cantidad END AS efecto
      FROM movimientos_inventario m
      JOIN almacenes a ON a.tipo = 'camion'
                      AND a.id IN (m.almacen_origen_id, m.almacen_destino_id)
      LEFT JOIN ventas v ON m.documento_tipo IN ('venta', 'venta_modificada')
                        AND v.id = m.documento_id
      LEFT JOIN mermas me ON m.documento_tipo = 'merma' AND me.id = m.documento_id
      LEFT JOIN dias k ON m.documento_tipo = 'carga' AND k.id = m.documento_id
)
SELECT camion, producto, dia, sum(efecto) AS cantidad
  FROM efectos
 WHERE dia >= :desde
 GROUP BY camion, producto, dia
"""

SQL_EN_LOS_CAMIONES = """
SELECT e.almacen_id AS camion, e.producto_id AS producto, e.cantidad
  FROM existencias e
  JOIN almacenes a ON a.id = e.almacen_id AND a.tipo = 'camion'
"""

# Con su precio de venta de hoy, para el camión en dinero (§185).
SQL_ARTICULOS = f"""
SELECT p.id AS producto, COALESCE(c.nombre, 'Sin familia') AS familia, c.orden,
       p.sku, p.nombre AS articulo, pv.precio
  FROM productos p
  LEFT JOIN categorias c ON c.id = p.categoria_id
  {precio_de_venta("p.id")}
 WHERE p.id = ANY(:ids)
"""

SQL_MAYOREO = """
SELECT m.fecha, m.folio, c.nombre_comercial AS cliente, u.nombre AS vendio, m.total
  FROM ventas_mayoreo m
  LEFT JOIN clientes c ON c.id = m.cliente_id
  LEFT JOIN usuarios u ON u.id = m.vendida_por
 WHERE m.fecha BETWEEN :desde AND :hasta AND m.estado = 'confirmada'
 ORDER BY m.fecha, m.folio
"""

# El corte, en el último día que contó (§181): el del lunes, con la carga «para
# el domingo», salía el domingo. El que todavía no se cierra, en el de su carga.
SQL_CORTES = """
SELECT COALESCE(l.cuenta_hasta, l.fecha_operativa) AS fecha, u.nombre AS vendedor,
       l.folio, l.estado,
       l.efectivo_esperado AS esperado, l.efectivo_entregado AS entregado,
       l.diferencia_efectivo AS diferencia, NULL::numeric AS declarado
  FROM liquidaciones l
  JOIN usuarios u ON u.id = l.vendedor_id
 WHERE COALESCE(l.cuenta_hasta, l.fecha_operativa) BETWEEN :desde AND :hasta
   AND l.estado <> 'anulada'
UNION ALL
-- El que mandó el vendedor y la oficina todavía no cierra.
SELECT k.fecha_operativa, u.nombre, NULL, 'por_cerrar', NULL, NULL, NULL,
       k.efectivo_declarado
  FROM cortes_vendedor k
  JOIN usuarios u ON u.id = k.vendedor_id
 WHERE k.fecha_operativa BETWEEN :desde AND :hasta AND k.estado = 'pendiente'
 ORDER BY 1, 2
"""

SQL_INVENTARIO = f"""
-- El camión con el nombre de quien lo trae: «Camión 01» solo no dice de quién es.
SELECT a.nombre || CASE WHEN a.tipo = 'camion' THEN COALESCE(' · ' || r.nombre, '') ELSE '' END
         AS almacen,
       a.tipo AS tipo_almacen,
       COALESCE(c.nombre, 'Sin familia') AS familia, p.sku, p.nombre AS articulo,
       e.cantidad AS piezas, pv.precio,
       round(e.cantidad * pv.precio, 2) AS valor
  FROM existencias e
  JOIN almacenes a ON a.id = e.almacen_id AND a.activo
  LEFT JOIN usuarios r ON r.id = a.responsable_id
  JOIN productos p ON p.id = e.producto_id
  LEFT JOIN categorias c ON c.id = p.categoria_id
  {precio_de_venta("p.id")}
 WHERE e.cantidad <> 0 AND a.tipo IN ('bodega', 'camion')
 ORDER BY (a.tipo <> 'bodega'), a.nombre, c.orden NULLS LAST, familia, p.nombre
"""

SQL_CARGAS = """
SELECT g.fecha_operativa AS fecha, g.folio, u.nombre AS vendedor, d.nombre AS camion,
       o.nombre AS bodega, g.estado, count(x.id) AS productos,
       COALESCE(sum(x.cantidad), 0) AS piezas
  FROM cargas g
  JOIN almacenes d ON d.id = g.almacen_destino_id
  JOIN almacenes o ON o.id = g.almacen_origen_id
  LEFT JOIN usuarios u ON u.id = g.vendedor_id
  LEFT JOIN carga_detalle x ON x.carga_id = g.id
 WHERE g.fecha_operativa BETWEEN :desde AND :hasta
   AND g.estado NOT IN ('borrador', 'cancelada')
 GROUP BY g.id, g.fecha_operativa, g.folio, u.nombre, d.nombre, o.nombre, g.estado
 ORDER BY g.fecha_operativa, g.folio
"""

SQL_MERMAS = """
SELECT m.fecha_operativa AS fecha, m.folio_local AS folio, u.nombre AS quien,
       CASE m.tipo WHEN 'devolucion_cliente' THEN 'Devolución del cliente'
                   WHEN 'cambio' THEN 'Cambio físico'
                   ELSE 'Merma del camión' END AS tipo,
       COALESCE(mm.nombre, m.motivo_codigo) AS motivo, p.nombre AS articulo,
       d.cantidad_base AS piezas
  FROM mermas m
  JOIN merma_detalle d ON d.merma_id = m.id
  JOIN productos p ON p.id = d.producto_id
  JOIN usuarios u ON u.id = m.vendedor_id
  LEFT JOIN motivos_merma mm ON mm.codigo = m.motivo_codigo
 WHERE m.fecha_operativa BETWEEN :desde AND :hasta AND m.estado = 'confirmada'
UNION ALL
SELECT s.fecha_operativa, s.folio, a.nombre, 'Merma de bodega',
       COALESCE(mm.nombre, s.motivo_codigo), p.nombre, x.cantidad
  FROM salidas s
  JOIN salida_detalle x ON x.salida_id = s.id
  JOIN productos p ON p.id = x.producto_id
  JOIN almacenes a ON a.id = s.almacen_origen_id
  LEFT JOIN motivos_merma mm ON mm.codigo = s.motivo_codigo
 WHERE s.tipo = 'merma' AND s.estado = 'confirmada'
   AND s.fecha_operativa BETWEEN :desde AND :hasta
 ORDER BY 1, 2, 6
"""

SQL_SALDOS = """
SELECT u.nombre AS vendedor, u.activo,
       COALESCE(c.cargos, 0) AS cargos, COALESCE(c.abonos, 0) AS abonos,
       COALESCE(c.condonado, 0) AS condonado, COALESCE(c.saldo, 0) AS saldo
  FROM usuarios u
  LEFT JOIN v_cuenta_vendedor c ON c.vendedor_id = u.id
 WHERE u.rol_codigo = 'vendedor' AND (u.activo OR COALESCE(c.saldo, 0) <> 0)
 ORDER BY COALESCE(c.saldo, 0) DESC, u.nombre
"""

# «Cuánto hubo en total de dinero en todo el producto cada fecha», «por camión
# y por almacén, y la suma» (§183). A precio de venta de hoy, como la hoja de
# Inventario: así lo que cambia de un día a otro es mercancía que entró o salió,
# no precios. Lo de hoy, y lo que se movió cada día, valuado igual; de ahí sale
# cuánto valía cada almacén al cerrar cada día.
_PRECIO_DE_HOY = f"""
    SELECT p.id AS producto_id, pv.precio
      FROM productos p
      {precio_de_venta("p.id")}
"""

SQL_VALOR_DE_LOS_ALMACENES = f"""
WITH precio AS ({_PRECIO_DE_HOY})
SELECT a.id, a.nombre, a.tipo, r.nombre AS vendedor,
       COALESCE(sum(e.cantidad * pr.precio), 0) AS valor
  FROM almacenes a
  LEFT JOIN usuarios r ON r.id = a.responsable_id
  LEFT JOIN existencias e ON e.almacen_id = a.id
  LEFT JOIN precio pr ON pr.producto_id = e.producto_id
 WHERE a.tipo IN ('bodega', 'camion')
 GROUP BY a.id, a.nombre, a.tipo, r.nombre
"""

# El día de cada movimiento es el de su documento si es venta, mayoreo o merma, y
# si no el día en que se hizo: la carga sale de la bodega la tarde en que se
# acepta, y al cerrar ese día ya está arriba del camión.
#
# Y qué fue (§184), para ver día por día que todo cuadra: lo que entró a la
# bodega, lo que vendieron los vendedores, el mayoreo, las mermas, las
# devoluciones de clientes, los ajustes de la oficina y lo eliminado en
# Operaciones. Lo que solo se movió de un almacén a otro —la carga, el traspaso,
# lo que regresó a la bodega— es `entre_almacenes`: cambia cada almacén, pero no
# la suma.
SQL_VALOR_MOVIDO = f"""
WITH precio AS ({_PRECIO_DE_HOY}),
efectos AS (
    SELECT a.id AS almacen,
           COALESCE(v.fecha_operativa, me.fecha_operativa, vm.fecha,
                    CAST(m.fecha_servidor AS date)) AS dia,
           CASE
             WHEN ao.tipo IN ('bodega', 'camion') AND ad.tipo IN ('bodega', 'camion')
               THEN 'entre_almacenes'
             WHEN m.documento_tipo IN ('venta', 'venta_modificada', 'venta_cancelada',
                                       'venta_corregida') THEN 'vendedores'
             WHEN m.documento_tipo = 'venta_mayoreo' THEN 'mayoreo'
             WHEN m.documento_tipo = 'merma' AND m.tipo = 'devolucion' THEN 'devoluciones'
             WHEN m.documento_tipo = 'merma' OR m.tipo = 'merma' THEN 'mermas'
             WHEN m.documento_tipo IN ('entrada', 'traspaso') OR m.tipo = 'compra'
               THEN 'entradas'
             WHEN m.documento_tipo = 'anulacion' THEN 'eliminado'
             ELSE 'ajustes'
           END AS categoria,
           CASE WHEN m.almacen_destino_id = a.id THEN m.cantidad
                ELSE -m.cantidad END * pr.precio AS valor
      FROM movimientos_inventario m
      JOIN almacenes a ON a.tipo IN ('bodega', 'camion')
                      AND a.id IN (m.almacen_origen_id, m.almacen_destino_id)
      JOIN precio pr ON pr.producto_id = m.producto_id AND pr.precio IS NOT NULL
      LEFT JOIN almacenes ao ON ao.id = m.almacen_origen_id
      LEFT JOIN almacenes ad ON ad.id = m.almacen_destino_id
      LEFT JOIN ventas v ON m.documento_tipo IN ('venta', 'venta_modificada')
                        AND v.id = m.documento_id
      LEFT JOIN mermas me ON m.documento_tipo = 'merma' AND me.id = m.documento_id
      LEFT JOIN ventas_mayoreo vm ON m.documento_tipo = 'venta_mayoreo'
                                 AND vm.id = m.documento_id
)
SELECT almacen, dia, categoria, sum(valor) AS valor
  FROM efectos
 WHERE dia >= :desde
 GROUP BY almacen, dia, categoria
"""

# Lo que cambió la mercancía de toda la empresa cada día, en el orden de la tabla.
CAMBIOS_DEL_DIA = {
    "entradas": "Entró a la bodega",
    "vendedores": "Vendieron los vendedores",
    "mayoreo": "Venta de mayoreo",
    "mermas": "Mermas",
    "devoluciones": "Devoluciones de clientes",
    "ajustes": "Ajustes y conteos",
    "eliminado": "Eliminado en Operaciones",
}

SQL_MOVIMIENTOS = """
SELECT m.fecha, u.nombre AS vendedor, m.tipo, m.origen, m.concepto, m.importe
  FROM cuenta_vendedor m
  JOIN usuarios u ON u.id = m.vendedor_id
 WHERE m.fecha BETWEEN :desde AND :hasta
 ORDER BY m.fecha, u.nombre, m.registrado_en
"""


async def _camiones(filas, rango) -> list[dict]:
    """Cada camión, día por día y artículo por artículo (§179).

    «Quedó» sale de lo que el camión tiene hoy, quitándole lo que se movió
    después de ese día; «Traía» es lo que quedó el día anterior. «Otros» es lo
    que no es carga, venta ni merma, así que la cuenta de cada día cierra
    siempre: traía + cargó − vendió − merma + otros = quedó. Solo salen los días
    en que el camión se movió.
    """
    desde, hasta = rango["desde"], rango["hasta"]

    async def por_dia(sql: str, parametros: dict) -> dict[tuple, Decimal]:
        juntos: dict[tuple, Decimal] = defaultdict(Decimal)
        for f in await filas(sql, parametros):
            juntos[(f["camion"], f["producto"], f["dia"])] += Decimal(f["cantidad"] or 0)
        return juntos

    cargado = await por_dia(SQL_CARGADO_AL_CAMION, rango)
    vendido = await por_dia(SQL_VENDIDO_DEL_CAMION, rango)
    merma = await por_dia(SQL_MERMA_DEL_CAMION, rango)
    movido = await por_dia(SQL_MOVIMIENTO_DEL_CAMION, rango)
    ahora = {
        (f["camion"], f["producto"]): Decimal(f["cantidad"])
        for f in await filas(SQL_EN_LOS_CAMIONES)
    }

    dias_del_camion: dict = defaultdict(set)
    for juntos in (cargado, vendido, merma, movido):
        for (camion, _, dia), cantidad in juntos.items():
            if desde <= dia <= hasta and cantidad:
                dias_del_camion[camion].add(dia)
    if not dias_del_camion:
        return []

    productos_del_camion: dict = defaultdict(set)
    for juntos in (cargado, vendido, merma, movido, ahora):
        for clave in juntos:
            if clave[0] in dias_del_camion:
                productos_del_camion[clave[0]].add(clave[1])
    articulos = {
        f["producto"]: f
        for f in await filas(
            SQL_ARTICULOS,
            {"ids": list({p for ps in productos_del_camion.values() for p in ps})},
        )
    }
    nombres = {f["camion"]: f for f in await filas(SQL_CAMIONES)}
    # Lo que se movió después del periodo, para llegar de hoy a su último día.
    despues: dict[tuple, Decimal] = defaultdict(Decimal)
    for (camion, producto, dia), cantidad in movido.items():
        if dia > hasta:
            despues[(camion, producto)] += cantidad

    camiones = []
    for camion, dias in dias_del_camion.items():
        dias = sorted(dias)
        renglones = []
        for producto in productos_del_camion[camion]:
            quedo = ahora.get((camion, producto), Decimal(0)) - despues[(camion, producto)]
            por_dia_del_articulo = {}
            dia = hasta
            while dia >= desde:
                clave = (camion, producto, dia)
                traia = quedo - movido.get(clave, Decimal(0))
                if dia in dias:
                    cargo = cargado.get(clave, Decimal(0))
                    vendio = vendido.get(clave, Decimal(0))
                    mermo = merma.get(clave, Decimal(0))
                    por_dia_del_articulo[dia] = {
                        "traia": traia, "cargo": cargo, "vendio": vendio,
                        "merma": mermo, "otros": quedo - traia - cargo + vendio + mermo,
                        "quedo": quedo,
                    }
                quedo = traia
                dia -= timedelta(days=1)
            if not any(any(v.values()) for v in por_dia_del_articulo.values()):
                continue
            a = articulos[producto]
            renglones.append({"familia": a["familia"], "sku": a["sku"],
                              "articulo": a["articulo"], "orden": a["orden"],
                              "precio": a["precio"],
                              "dias": por_dia_del_articulo})
        renglones.sort(key=lambda r: (r["orden"] is None, r["orden"] or 0, r["familia"],
                                      r["articulo"]))
        nombre = nombres.get(camion) or {"nombre": "Camión", "vendedor": None}
        camiones.append({"camion": nombre["nombre"], "vendedor": nombre["vendedor"],
                         "dias": dias, "articulos": renglones})
    camiones.sort(key=lambda c: (c["vendedor"] is None, c["vendedor"] or "", c["camion"]))
    return camiones


async def _valor_por_dia(filas, rango) -> dict:
    """Lo que valía cada almacén al cerrar cada día del periodo (§183), y en qué
    se fue o de dónde vino lo que cambió cada día (§184).

    Se parte de lo que vale hoy y se va hacia atrás quitando lo que se movió
    cada día, como el «Quedó» del camión. Hasta hoy: de mañana no se sabe nada.
    Solo salen los almacenes que tuvieron mercancía en el periodo.
    """
    hoy = (await filas("SELECT CURRENT_DATE AS hoy"))[0]["hoy"]
    desde, ultimo = rango["desde"], min(rango["hasta"], hoy)
    if desde > ultimo:
        return {"almacenes": [], "dias": []}
    almacenes = await filas(SQL_VALOR_DE_LOS_ALMACENES)
    valor = {a["id"]: Decimal(a["valor"]) for a in almacenes}
    movido: dict[tuple, Decimal] = defaultdict(Decimal)
    por_que: dict[tuple, Decimal] = defaultdict(Decimal)
    for f in await filas(SQL_VALOR_MOVIDO, {"desde": desde}):
        if f["almacen"] not in valor:
            continue
        cuanto = Decimal(f["valor"] or 0)
        movido[(f["almacen"], f["dia"])] += cuanto
        por_que[(f["dia"], f["categoria"])] += cuanto
    for (almacen, dia), cuanto in movido.items():
        if dia > ultimo:
            valor[almacen] -= cuanto

    dias = []
    dia = ultimo
    while dia >= desde:
        cerro = sum(valor.values(), Decimal(0))
        cambios = {c: por_que.get((dia, c), Decimal(0)) for c in CAMBIOS_DEL_DIA}
        dias.append({"dia": dia, "valores": dict(valor), "cerro": cerro,
                     "cambios": cambios,
                     # Lo que solo cambió de almacén no cambia la suma.
                     "empezo": cerro - sum(cambios.values(), Decimal(0))})
        for almacen in valor:
            valor[almacen] -= movido.get((almacen, dia), Decimal(0))
        dia -= timedelta(days=1)
    dias.reverse()

    con_mercancia = [a for a in almacenes if any(d["valores"][a["id"]] for d in dias)]
    # La bodega primero y los camiones por quien los trae, como en Inventario.
    con_mercancia.sort(key=lambda a: (a["tipo"] != "bodega", a["vendedor"] is None,
                                      a["vendedor"] or "", a["nombre"]))
    return {"almacenes": con_mercancia, "dias": dias}


async def datos_del_reporte(sesion, periodo) -> Reporte:
    """Todo lo que lleva el reporte, del periodo que se eligió."""
    rango = {"desde": periodo.inicio, "hasta": periodo.fin}

    async def filas(sql: str, parametros: dict | None = None) -> list[dict]:
        resultado = await sesion.execute(text(sql), parametros or {})
        return [dict(f) for f in resultado.mappings().all()]

    return Reporte(
        desde=periodo.inicio,
        hasta=periodo.fin,
        descripcion=periodo.descripcion,
        ventas=await filas(SQL_VENTAS, rango),
        camiones=await _camiones(filas, rango),
        mayoreo=await filas(SQL_MAYOREO, rango),
        cortes=await filas(SQL_CORTES, rango),
        inventario=await filas(SQL_INVENTARIO),
        cargas=await filas(SQL_CARGAS, rango),
        mermas=await filas(SQL_MERMAS, rango),
        saldos=await filas(SQL_SALDOS),
        movimientos=await filas(SQL_MOVIMIENTOS, rango),
        valor_por_dia=await _valor_por_dia(filas, rango),
    )


# ---------------------------------------------------------------------------
# Las palabras
# ---------------------------------------------------------------------------


def _pesos(valor) -> str:
    return f"${Decimal(valor or 0):,.2f}"


def resultado_del_corte(c: dict) -> str:
    """Lo que dice la columna «Cómo quedó», en palabras."""
    if c["estado"] == "por_cerrar":
        return f"Falta cerrarlo: el vendedor dice que entrega {_pesos(c['declarado'])}"
    if c["estado"] != "cerrada":
        return "Abierto: todavía no se cierra"
    diferencia = Decimal(c["diferencia"] or 0)
    if diferencia == 0:
        return "Cuadró"
    if diferencia < 0:
        return f"Faltaron {_pesos(-diferencia)}"
    return f"Sobraron {_pesos(diferencia)}"


def _con_dia(fecha) -> str:
    return f"{_DIAS[fecha.weekday()]} {fecha:%d/%m/%Y}"


# ---------------------------------------------------------------------------
# El libro
# ---------------------------------------------------------------------------


class _Hoja:
    """Una hoja que se escribe de arriba abajo: título, tablas y totales."""

    def __init__(self, libro: Workbook, nombre: str, *, primera: bool = False):
        self.hoja = libro.active if primera else libro.create_sheet()
        self.hoja.title = nombre
        self.fila = 1
        self.anchos: dict[int, int] = {}

    def titulo(self, texto: str, subtitulo: str = "") -> None:
        celda = self.hoja.cell(self.fila, 1, texto)
        celda.font = Font(bold=True, size=16, color=ROJO)
        self.fila += 1
        if subtitulo:
            self.hoja.cell(self.fila, 1, subtitulo).font = Font(italic=True, color="666666")
            self.fila += 1
        self.fila += 1

    def nota(self, texto: str) -> None:
        self.hoja.cell(self.fila, 1, texto).font = Font(italic=True, color="666666")
        self.fila += 1

    def tabla(
        self,
        titulo: str,
        columnas: list[tuple[str, str]],
        renglones: list[list],
        *,
        totales: list | None = None,
        vacio: str = "No hubo en este periodo.",
    ) -> None:
        """`columnas` es (encabezado, formato): «texto», «fecha», «entero» o
        «dinero». `totales`, el renglón final en negritas (None en lo que no se
        suma)."""
        self.hoja.cell(self.fila, 1, titulo).font = Font(bold=True, size=12)
        self.fila += 1
        if not renglones:
            self.nota(vacio)
            self.fila += 1
            return
        for n, (encabezado, _) in enumerate(columnas, start=1):
            celda = self.hoja.cell(self.fila, n, encabezado)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor=ROJO)
            celda.alignment = Alignment(vertical="center", wrap_text=True)
            self._ancho(n, encabezado)
        self.fila += 1
        for renglon in renglones:
            self._renglon(columnas, renglon)
        if totales is not None:
            self._renglon(columnas, totales, negritas=True)
        self.fila += 1

    def _renglon(self, columnas, valores, *, negritas: bool = False) -> None:
        for n, ((_, formato), valor) in enumerate(zip(columnas, valores, strict=True), start=1):
            if isinstance(valor, Decimal):
                valor = float(valor)
            celda = self.hoja.cell(self.fila, n, valor)
            if formato == "dinero":
                celda.number_format = DINERO
            elif formato == "entero":
                celda.number_format = ENTERO
            elif formato == "fecha":
                celda.number_format = FECHA
            if negritas:
                celda.font = Font(bold=True)
                celda.fill = PatternFill("solid", fgColor=GRIS)
            self._ancho(n, _como_texto(valor, formato))
        self.fila += 1

    def _ancho(self, columna: int, texto: str) -> None:
        self.anchos[columna] = max(self.anchos.get(columna, 8), min(len(texto) + 2, 60))

    def terminar(self) -> None:
        for columna, ancho in self.anchos.items():
            self.hoja.column_dimensions[get_column_letter(columna)].width = ancho
        # Para imprimirla: de lado y a lo ancho de una hoja, que una tabla partida
        # en dos páginas no se entiende.
        self.hoja.page_setup.orientation = "landscape"
        self.hoja.sheet_properties.pageSetUpPr.fitToPage = True
        self.hoja.page_setup.fitToWidth = 1
        self.hoja.page_setup.fitToHeight = 0


def _como_texto(valor, formato: str) -> str:
    if valor is None:
        return ""
    if isinstance(valor, str):
        return valor
    if formato == "dinero":
        return f"${float(valor):,.2f}"
    if formato == "fecha":
        return "00/00/0000"
    if formato == "entero":
        return f"{float(valor):,.0f}"
    return str(valor)


def _suma(renglones: list[dict], clave: str) -> Decimal:
    return sum((Decimal(r[clave] or 0) for r in renglones), Decimal(0))


DEL_DIA = (("Traía", "traia"), ("Cargó", "cargo"), ("Vendió", "vendio"),
           ("Merma", "merma"), ("Otros", "otros"), ("Quedó", "quedo"))
AMARILLO = "FFF2CC"


def _nombre_de_hoja(libro: Workbook, nombre: str) -> str:
    """Lo que Excel deja: sin []:*?/\\, hasta 31 letras y sin repetir."""
    limpio = "".join(c for c in nombre if c not in "[]:*?/\\").strip()[:31] or "Camión"
    propuesto, n = limpio, 2
    while propuesto in libro.sheetnames:
        sufijo = f" ({n})"
        propuesto, n = limpio[: 31 - len(sufijo)] + sufijo, n + 1
    return propuesto


def _hoja_del_camion(libro: Workbook, camion: dict, periodo: str) -> None:
    """Los artículos hacia abajo y, por cada día, lo que traía, cargó, vendió,
    mermó, lo demás y lo que le quedó (§179)."""
    hoja = libro.create_sheet(_nombre_de_hoja(libro, camion["vendedor"] or camion["camion"]))
    quien = camion["vendedor"] or camion["camion"]
    hoja.cell(1, 1, f"El camión de {quien}, día por día").font = Font(
        bold=True, size=16, color=ROJO
    )
    hoja.cell(2, 1, f"{camion['camion']} · {periodo}").font = Font(italic=True,
                                                                  color="666666")
    dias = camion["dias"]
    arriba, abajo = 4, 5
    blanco = Font(bold=True, color="FFFFFF")
    for n, encabezado in enumerate(("Familia", "Clave", "Artículo"), start=1):
        hoja.merge_cells(start_row=arriba, start_column=n, end_row=abajo, end_column=n)
        celda = hoja.cell(arriba, n, encabezado)
        celda.font, celda.fill = blanco, PatternFill("solid", fgColor=ROJO)
        celda.alignment = Alignment(vertical="center")
    for i, dia in enumerate(dias):
        primera = 4 + i * len(DEL_DIA)
        hoja.merge_cells(start_row=arriba, start_column=primera,
                         end_row=arriba, end_column=primera + len(DEL_DIA) - 1)
        celda = hoja.cell(arriba, primera, _con_dia(dia))
        celda.font = blanco
        celda.fill = PatternFill("solid", fgColor=ROJO if i % 2 == 0 else "A61E22")
        celda.alignment = Alignment(horizontal="center")
        for j, (encabezado, _) in enumerate(DEL_DIA):
            celda = hoja.cell(abajo, primera + j, encabezado)
            celda.font = Font(bold=True)
            celda.fill = PatternFill("solid", fgColor=GRIS)
            celda.alignment = Alignment(horizontal="center")

    fila = abajo + 1
    totales = {(dia, clave): Decimal(0) for dia in dias for _, clave in DEL_DIA}
    for r in camion["articulos"]:
        for n, valor in enumerate((r["familia"], r["sku"], r["articulo"]), start=1):
            hoja.cell(fila, n, valor)
        for i, dia in enumerate(dias):
            cifras = r["dias"][dia]
            for j, (_, clave) in enumerate(DEL_DIA):
                celda = hoja.cell(fila, 4 + i * len(DEL_DIA) + j, float(cifras[clave]))
                celda.number_format = ENTERO
                if clave == "quedo":
                    celda.font = Font(bold=True)
                # Lo que el camión ganó o perdió sin carga, venta ni merma: que se vea.
                if clave == "otros" and cifras[clave]:
                    celda.fill = PatternFill("solid", fgColor=AMARILLO)
                totales[(dia, clave)] += cifras[clave]
        fila += 1
    hoja.cell(fila, 1, "Total").font = Font(bold=True)
    for n in range(1, 4 + len(dias) * len(DEL_DIA)):
        hoja.cell(fila, n).fill = PatternFill("solid", fgColor=GRIS)
    for i, dia in enumerate(dias):
        for j, (_, clave) in enumerate(DEL_DIA):
            celda = hoja.cell(fila, 4 + i * len(DEL_DIA) + j, float(totales[(dia, clave)]))
            celda.number_format, celda.font = ENTERO, Font(bold=True)
    fila += 1

    # «Que el Excel muestre el valor del carro del vendedor cada día» (§185): el
    # mismo total, en pesos, a precio de venta de hoy como en Inventario.
    hoja.cell(fila, 1, "En dinero").font = Font(bold=True)
    for n in range(1, 4 + len(dias) * len(DEL_DIA)):
        hoja.cell(fila, n).fill = PatternFill("solid", fgColor=GRIS)
    con_precio = [r for r in camion["articulos"] if r["precio"] is not None]
    for i, dia in enumerate(dias):
        for j, (_, clave) in enumerate(DEL_DIA):
            pesos = sum((r["dias"][dia][clave] * Decimal(r["precio"]) for r in con_precio),
                        Decimal(0))
            celda = hoja.cell(fila, 4 + i * len(DEL_DIA) + j, float(pesos))
            celda.number_format, celda.font = PESOS, Font(bold=True)
    fila += 2
    sin_precio = [r["articulo"] for r in camion["articulos"] if r["precio"] is None]
    for texto in (
        "Piezas sueltas. Traía: con lo que empezó el día. Quedó: con lo que lo terminó.",
        "En dinero: lo mismo a precio de venta de hoy. Traía más cargó es con lo que "
        "salió a vender; lo que cobró de verdad está en Ventas y efectivo.",
        "La carga sale el día en que se vendió con ella; la merma es la del camión "
        "(merma y cambio).",
        "Otros, en amarillo: lo que entró o salió sin ser carga, venta ni merma —ajustes "
        "de la oficina, devoluciones de clientes, lo que regresó a la bodega o una "
        "venta eliminada—.",
        "Solo salen los días en que el camión se movió.",
        *([f"Sin precio, no suman en dinero: {', '.join(sin_precio)}."] if sin_precio else []),
    ):
        hoja.cell(fila, 1, texto).font = Font(italic=True, color="666666")
        fila += 1

    hoja.column_dimensions["A"].width = 16
    hoja.column_dimensions["B"].width = 12
    hoja.column_dimensions["C"].width = 32
    # Lo bastante ancho para «$12,345» del renglón en dinero.
    for n in range(4, 4 + len(dias) * len(DEL_DIA)):
        hoja.column_dimensions[get_column_letter(n)].width = 9
    # Al moverse a la derecha, el artículo y los días se quedan a la vista.
    hoja.freeze_panes = hoja.cell(abajo + 1, 4)
    hoja.page_setup.orientation = "landscape"
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0


def armar_libro(reporte: Reporte, *, generado_por: str, generado_en: datetime) -> bytes:
    """El archivo .xlsx, listo para mandarlo."""
    libro = Workbook()
    periodo = reporte.descripcion
    if periodo and periodo[0].islower():
        periodo = periodo[0].upper() + periodo[1:]

    ventas_importe = _suma(reporte.ventas, "importe")
    ventas_efectivo = _suma(reporte.ventas, "efectivo")
    separar_efectivo = ventas_efectivo != ventas_importe
    cuantas_ventas = int(sum(r["ventas"] for r in reporte.ventas))
    mayoreo_importe = _suma(reporte.mayoreo, "total")
    cerrados = [c for c in reporte.cortes if c["estado"] == "cerrada"]
    piezas_cargadas = _suma(reporte.cargas, "piezas")
    piezas_mermadas = _suma(reporte.mermas, "piezas")
    bodega = [r for r in reporte.inventario if r["tipo_almacen"] == "bodega"]
    camiones = [r for r in reporte.inventario if r["tipo_almacen"] == "camion"]
    sin_precio = sorted({r["articulo"] for r in reporte.inventario if r["precio"] is None})
    deben = _suma([s for s in reporte.saldos if Decimal(s["saldo"]) > 0], "saldo")

    # --- Resumen ------------------------------------------------------------
    resumen = _Hoja(libro, "Resumen", primera=True)
    resumen.titulo(
        "Reporte de Distribuciones SE",
        f"{periodo} · hecho el {generado_en:%d/%m/%Y a las %H:%M} por {generado_por}",
    )
    resumen.tabla(
        "Lo más importante",
        [("Qué", "texto"), ("Cuánto", "dinero"), ("Para entenderlo", "texto")],
        [
            *(
                [
                    ["Ventas de los vendedores", ventas_importe,
                     f"{cuantas_ventas:,} ventas en la calle"],
                    ["De eso, en efectivo", ventas_efectivo,
                     "Lo que los vendedores tienen que entregar en el corte"],
                ]
                if separar_efectivo
                else [
                    ["Ventas de los vendedores", ventas_importe,
                     f"{cuantas_ventas:,} ventas en la calle, todas en efectivo: es lo "
                     "que tienen que entregar en el corte"],
                ]
            ),
            ["Efectivo que entregaron", _suma(cerrados, "entregado"),
             f"En {len(cerrados)} corte(s) ya cerrados"],
            ["Diferencia en los cortes", _suma(cerrados, "diferencia"),
             "Negativo es lo que faltó; positivo, lo que sobró"],
            ["Ventas de mayoreo (bodega)", mayoreo_importe,
             f"{len(reporte.mayoreo)} venta(s) de la oficina"],
            ["Inventario en la bodega (hoy)", _suma(bodega, "valor"),
             "A precio de venta"],
            ["Inventario en los camiones (hoy)", _suma(camiones, "valor"),
             "A precio de venta"],
            ["Lo que deben los vendedores (hoy)", deben,
             "Faltantes y mermas a su cargo que no han pagado"],
        ],
    )
    resumen.tabla(
        "Mercancía",
        [("Qué", "texto"), ("Piezas", "entero"), ("Para entenderlo", "texto")],
        [
            ["Cargas a los camiones", piezas_cargadas, f"{len(reporte.cargas)} carga(s)"],
            ["Mermas y devoluciones", piezas_mermadas,
             f"{len({(m['folio'], m['fecha']) for m in reporte.mermas})} registro(s)"],
        ],
    )
    resumen.nota("Cada tema tiene su hoja abajo: Ventas y efectivo, una por cada "
                 "vendedor con el inventario de su camión día por día, Inventario, "
                 "Día por día —el dinero y la mercancía de cada día, para ver que todo "
                 "cuadra—, Cargas y mermas, y Cuenta de vendedores.")
    if sin_precio:
        resumen.nota(f"Ojo: {len(sin_precio)} artículo(s) sin precio de venta no suman "
                     f"en el inventario: {', '.join(sin_precio[:6])}"
                     + ("…" if len(sin_precio) > 6 else "."))
    resumen.terminar()

    # --- Ventas y efectivo -----------------------------------------------------
    hoja = _Hoja(libro, "Ventas y efectivo")
    hoja.titulo("Ventas y efectivo", periodo)
    # La columna «En efectivo» repetiría «Importe» peso por peso: sale solo si
    # algo del periodo no fue en efectivo.
    en_efectivo = [("En efectivo", "dinero")] if separar_efectivo else []

    def _y_efectivo(cifra) -> list:
        return [cifra] if separar_efectivo else []

    hoja.tabla(
        "Ventas de cada vendedor, día por día",
        [("Día", "texto"), ("Vendedor", "texto"), ("Ventas", "entero"),
         ("Piezas", "entero"), ("Importe", "dinero"), *en_efectivo],
        [[_con_dia(r["fecha"]), r["vendedor"], r["ventas"], r["piezas"], r["importe"],
          *_y_efectivo(r["efectivo"])] for r in reporte.ventas],
        totales=["Total", "", cuantas_ventas,
                 _suma(reporte.ventas, "piezas"), ventas_importe,
                 *_y_efectivo(ventas_efectivo)],
    )
    por_vendedor: dict[str, dict] = {}
    for r in reporte.ventas:
        v = por_vendedor.setdefault(
            r["vendedor"], {"ventas": 0, "importe": Decimal(0), "efectivo": Decimal(0)}
        )
        v["ventas"] += r["ventas"]
        v["importe"] += Decimal(r["importe"])
        v["efectivo"] += Decimal(r["efectivo"])
    hoja.tabla(
        "Total de cada vendedor",
        [("Vendedor", "texto"), ("Ventas", "entero"), ("Importe", "dinero"),
         *en_efectivo],
        [[n, v["ventas"], v["importe"], *_y_efectivo(v["efectivo"])]
         for n, v in sorted(por_vendedor.items(), key=lambda par: -par[1]["importe"])],
        totales=["Total", sum(v["ventas"] for v in por_vendedor.values()), ventas_importe,
                 *_y_efectivo(ventas_efectivo)],
    )
    if not separar_efectivo and reporte.ventas:
        hoja.nota("Todas las ventas del periodo fueron en efectivo.")
        hoja.fila += 1
    hoja.tabla(
        "Cortes: el efectivo que entregó cada vendedor",
        [("Día", "texto"), ("Vendedor", "texto"), ("Folio", "texto"),
         ("Debía entregar", "dinero"), ("Entregó", "dinero"), ("Diferencia", "dinero"),
         ("Cómo quedó", "texto")],
        [[_con_dia(c["fecha"]), c["vendedor"], c["folio"] or "",
          c["esperado"], c["entregado"] if c["estado"] == "cerrada" else None,
          c["diferencia"] if c["estado"] == "cerrada" else None, resultado_del_corte(c)]
         for c in reporte.cortes],
        totales=["Total (cerrados)", "", "", _suma(cerrados, "esperado"),
                 _suma(cerrados, "entregado"), _suma(cerrados, "diferencia"), ""],
        vacio="No hubo cortes en este periodo.",
    )
    hoja.tabla(
        "Ventas de mayoreo (de la bodega)",
        [("Fecha", "fecha"), ("Folio", "texto"), ("Cliente", "texto"),
         ("La hizo", "texto"), ("Importe", "dinero")],
        [[m["fecha"], m["folio"], m["cliente"] or "", m["vendio"] or "", m["total"]]
         for m in reporte.mayoreo],
        totales=["Total", "", "", "", mayoreo_importe],
    )
    hoja.terminar()

    # --- El camión de cada vendedor, día por día -----------------------------
    # «Quiero poder comparar por vendedor día a día, ver cómo se comporta su
    # inventario» (§179): una hoja por camión, los artículos hacia abajo y los
    # días hacia la derecha.
    if not reporte.camiones:
        hoja = _Hoja(libro, "Camiones")
        hoja.titulo("El camión de cada vendedor, día por día", periodo)
        hoja.nota("Los camiones no se movieron en este periodo.")
        hoja.terminar()
    for camion in reporte.camiones:
        _hoja_del_camion(libro, camion, periodo)

    # --- Inventario -------------------------------------------------------------
    hoja = _Hoja(libro, "Inventario")
    hoja.titulo("Inventario de hoy",
                f"Lo que hay al {generado_en:%d/%m/%Y %H:%M}, a precio de venta")
    almacenes = list(dict.fromkeys(r["almacen"] for r in reporte.inventario))
    if not almacenes:
        hoja.nota("No hay mercancía en ningún almacén.")
    for almacen in almacenes:
        suyos = [r for r in reporte.inventario if r["almacen"] == almacen]
        hoja.tabla(
            almacen,
            [("Familia", "texto"), ("Clave", "texto"), ("Artículo", "texto"),
             ("Piezas", "entero"), ("Precio", "dinero"), ("Vale", "dinero")],
            [[r["familia"], r["sku"], r["articulo"], r["piezas"],
              r["precio"] if r["precio"] is not None else "sin precio", r["valor"]]
             for r in suyos],
            totales=["Total", "", "", _suma(suyos, "piezas"), None, _suma(suyos, "valor")],
        )
    if sin_precio:
        hoja.nota("Los que no tienen precio no suman en «Vale»: " + ", ".join(sin_precio))
    hoja.terminar()

    # --- Día por día ----------------------------------------------------------------
    # «Quiero en el Excel ser capaz de distinguir que todo va bien día a día;
    # separa la venta a mayoreo de la de los vendedores» (§184). Y «cuánto hubo
    # de dinero en todo el producto cada fecha, por camión y por almacén, y la
    # suma» (§183).
    hoja = _Hoja(libro, "Día por día")
    hoja.titulo("Día por día: el dinero y la mercancía", periodo)
    valor = reporte.valor_por_dia
    dias = valor.get("dias", [])
    almacenes = valor.get("almacenes", [])

    vendido: dict = defaultdict(Decimal)
    for r in reporte.ventas:
        vendido[r["fecha"]] += Decimal(r["importe"])
    entregado: dict = defaultdict(Decimal)
    for c in cerrados:
        entregado[c["fecha"]] += Decimal(c["entregado"] or 0)
    de_mayoreo: dict = defaultdict(Decimal)
    for m in reporte.mayoreo:
        de_mayoreo[m["fecha"]] += Decimal(m["total"])
    dinero = [
        [_con_dia(d["dia"]), vendido[d["dia"]], entregado[d["dia"]],
         entregado[d["dia"]] - vendido[d["dia"]], de_mayoreo[d["dia"]],
         vendido[d["dia"]] + de_mayoreo[d["dia"]]]
        for d in dias
    ]
    hoja.tabla(
        "El dinero de cada día",
        [("Día", "texto"), ("Vendieron los vendedores", "dinero"),
         ("Entregaron en sus cortes", "dinero"), ("Diferencia", "dinero"),
         ("Venta de mayoreo", "dinero"), ("Total vendido", "dinero")],
        dinero,
        totales=["Total", *[sum((r[n] for r in dinero), Decimal(0)) for n in range(1, 6)]]
        if dinero else None,
        vacio="El periodo todavía no empieza.",
    )
    hoja.nota("Diferencia: lo entregado menos lo vendido por los vendedores. Negativo es "
              "lo que falta por entregar: un corte que no se ha cerrado, o uno que no "
              "cuadró. El mayoreo lo cobra la oficina y va aparte.")
    hoja.fila += 1

    mercancia = [
        [_con_dia(d["dia"]), d["empezo"], *d["cambios"].values(), d["cerro"]]
        for d in dias
    ]
    hoja.tabla(
        "La mercancía de cada día, a precio de venta",
        [("Día", "texto"), ("Valía al empezar", "dinero"),
         *[(nombre, "dinero") for nombre in CAMBIOS_DEL_DIA.values()],
         ("Valía al cerrar", "dinero")],
        mercancia,
        totales=["Total", dias[0]["empezo"],
                 *[sum((d["cambios"][c] for d in dias), Decimal(0)) for c in CAMBIOS_DEL_DIA],
                 dias[-1]["cerro"]] if dias else None,
        vacio="El periodo todavía no empieza.",
    )
    hoja.nota("Cada renglón se suma de izquierda a derecha: lo que valía al empezar, más lo "
              "que entró, menos lo que salió (va en negativo), es lo que valía al cerrar.")
    hoja.nota("Lo que solo se movió de un almacén a otro —la carga, lo que regresó a la "
              "bodega— no cambia la suma: se ve en la tabla de abajo.")
    hoja.fila += 1

    hoja.tabla(
        "Lo que valía cada almacén al cerrar el día",
        [("Día", "texto"),
         *[(a["nombre"] + (f" · {a['vendedor']}" if a["vendedor"] else ""), "dinero")
           for a in almacenes],
         ("Total", "dinero")],
        [[_con_dia(d["dia"]), *[d["valores"][a["id"]] for a in almacenes], d["cerro"]]
         for d in dias] if almacenes else [],
        vacio="No hubo mercancía en estos días, o el periodo todavía no empieza.",
    )
    hoja.nota("Al cerrar el día: el camión ya con la carga que se hizo en la tarde para el "
              "día siguiente, y sin lo que vendió.")
    hoja.nota("A los precios de hoy, como en Inventario: lo que cambia de un día a otro es "
              "mercancía que entró o salió, no precios. Por eso «Vendieron los vendedores» "
              "de la mercancía puede no ser lo mismo que lo que cobraron: el descuento o "
              "un precio que cambió.")
    if sin_precio:
        hoja.nota("Los que no tienen precio no suman: " + ", ".join(sin_precio))
    hoja.terminar()

    # --- Cargas y mermas --------------------------------------------------------
    hoja = _Hoja(libro, "Cargas y mermas")
    hoja.titulo("Cargas y mermas", periodo)
    hoja.tabla(
        "Cargas a los camiones",
        [("Fecha", "fecha"), ("Folio", "texto"), ("Vendedor", "texto"),
         ("Camión", "texto"), ("Salió de", "texto"), ("Artículos", "entero"),
         ("Piezas", "entero"), ("Cómo va", "texto")],
        [[c["fecha"], c["folio"], c["vendedor"] or "", c["camion"], c["bodega"],
          c["productos"], c["piezas"], ESTADO_DE_LA_CARGA.get(c["estado"], c["estado"])]
         for c in reporte.cargas],
        totales=["Total", "", "", "", "", None, piezas_cargadas, ""],
    )
    hoja.tabla(
        "Mermas y devoluciones",
        [("Fecha", "fecha"), ("Folio", "texto"), ("Quién / dónde", "texto"),
         ("Qué fue", "texto"), ("Motivo", "texto"), ("Artículo", "texto"),
         ("Piezas", "entero")],
        [[m["fecha"], m["folio"], m["quien"], m["tipo"], m["motivo"], m["articulo"],
          m["piezas"]] for m in reporte.mermas],
        totales=["Total", "", "", "", "", "", piezas_mermadas],
    )
    hoja.terminar()

    # --- Cuenta de vendedores --------------------------------------------------
    hoja = _Hoja(libro, "Cuenta de vendedores")
    hoja.titulo(
        "Cuenta de cada vendedor",
        "Lo que se le carga cuando le falta mercancía o efectivo, lo que paga y lo "
        "que debe",
    )
    hoja.tabla(
        f"Cómo está cada uno hoy ({generado_en:%d/%m/%Y})",
        [("Vendedor", "texto"), ("Se le ha cargado", "dinero"), ("Ha pagado", "dinero"),
         ("Se le perdonó", "dinero"), ("Debe", "dinero")],
        [[s["vendedor"] + ("" if s["activo"] else " (ya no trabaja)"), s["cargos"],
          s["abonos"], s["condonado"], s["saldo"]] for s in reporte.saldos],
        totales=["Total", _suma(reporte.saldos, "cargos"), _suma(reporte.saldos, "abonos"),
                 _suma(reporte.saldos, "condonado"), _suma(reporte.saldos, "saldo")],
        vacio="No hay vendedores dados de alta.",
    )
    hoja.tabla(
        f"Lo que se movió {reporte.descripcion}",
        [("Fecha", "fecha"), ("Vendedor", "texto"), ("Qué", "texto"),
         ("Por qué", "texto"), ("Detalle", "texto"), ("Importe", "dinero")],
        [[m["fecha"], m["vendedor"], TIPO_DE_MOVIMIENTO.get(m["tipo"], m["tipo"]),
          ORIGENES.get(m["origen"], m["origen"]), m["concepto"], m["importe"]]
         for m in reporte.movimientos],
    )
    hoja.terminar()

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def nombre_del_archivo(reporte: Reporte) -> str:
    if reporte.desde == reporte.hasta:
        return f"reporte_{reporte.desde:%Y-%m-%d}.xlsx"
    return f"reporte_{reporte.desde:%Y-%m-%d}_a_{reporte.hasta:%Y-%m-%d}.xlsx"
