"""El reporte en Excel para el administrador (ADR 0002 §127).

«Quiero generar un reporte en Excel para la persona que usa el perfil de
administrador; él no sabe usar mucho sistemas.» Con las cuatro cosas que eligió
—ventas y efectivo, inventario, cargas y mermas, y la cuenta de cada vendedor—
del periodo que se elija al bajarlo.

1. **Se baja del dashboard** como un .xlsx con su nombre, con una hoja de
   Resumen y una por tema, y dice las mismas cifras que el dashboard.
2. **Lo eliminado no cuenta** en ninguna hoja.
3. **Quién**: la oficina —quien ve el Tablero: administrador, gerencia y
   supervisor—; sin entrar al panel, nadie.
4. **Todo es en efectivo** (§170): lo vendido no se repite como «en efectivo»,
   salvo que en el periodo haya una transferencia de un teléfono viejo.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from tests.test_cierre_del_vendedor import _cerrar, _corte, dia  # noqa: F401
from tests.test_cortes_api import _cab, _usuario
from tests.test_operaciones import _anular
from tests.test_panel_liquidacion import _entrar

pytestmark = pytest.mark.asyncio

HOJAS = ["Resumen", "Ventas y efectivo", "Juan Pérez", "Inventario",
         "Día por día", "Cargas y mermas", "Cuenta de vendedores"]
DEL_DIA = ["Traía", "Cargó", "Vendió", "Merma", "Otros", "Quedó"]


async def _bajar(cliente, periodo="hoy"):
    r = await cliente.get(f"/panel/reporte/excel?periodo={periodo}")
    assert r.status_code == 200, r.text
    return r, load_workbook(io.BytesIO(r.content))


def _renglon(hoja, primera: str) -> list:
    """Los valores del renglón cuya primera columna dice `primera`."""
    for fila in hoja.iter_rows(values_only=True):
        if fila and fila[0] == primera:
            return list(fila)
    raise AssertionError(f"no está «{primera}» en {hoja.title}")


def _con(hoja, texto: str) -> list:
    """El primer renglón que tiene `texto` en alguna celda."""
    for fila in hoja.iter_rows(values_only=True):
        if any(c == texto for c in fila):
            return list(fila)
    raise AssertionError(f"no está «{texto}» en {hoja.title}")


async def test_el_reporte_trae_sus_hojas_con_las_cifras_del_dashboard(
    cliente, sesion, semilla, dia  # noqa: F811
):
    # El vendedor declara $2,200 de $2,250 y el gerente lo cierra: faltan $50.
    corte_id = await _corte(sesion, dia, semilla, efectivo="2200.00")
    await _cerrar(cliente, corte_id, await _cab(cliente, "GER01"))

    await _entrar(cliente)
    pagina = await cliente.get("/panel/reporte?periodo=hoy")
    assert pagina.status_code == 200 and 'id="bajar_reporte"' in pagina.text
    tablero = await cliente.get("/panel")
    assert 'id="ir_al_reporte"' in tablero.text

    r, libro = await _bajar(cliente)
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    hoy = date.today()
    assert f'filename="reporte_{hoy:%Y-%m-%d}.xlsx"' in r.headers["content-disposition"]
    assert libro.sheetnames == HOJAS

    resumen = libro["Resumen"]
    assert resumen["A1"].value == "Reporte de Distribuciones SE"
    assert "hecho el" in resumen["A2"].value and "por Bryan" in resumen["A2"].value
    assert _renglon(resumen, "Ventas de los vendedores")[1] == 2250
    assert _renglon(resumen, "Efectivo que entregaron")[1] == 2200
    assert _renglon(resumen, "Diferencia en los cortes")[1] == -50
    assert _renglon(resumen, "Lo que deben los vendedores (hoy)")[1] == 50
    assert _renglon(resumen, "Cargas a los camiones")[1] == 240

    ventas = libro["Ventas y efectivo"]
    fila = _con(ventas, "Juan Pérez")
    assert fila[2:5] == [1, 180, 2250]

    # El camión de cada vendedor, día por día (§177, §179): no traía nada, cargó
    # 240, vendió 180 y le quedaron 60.
    from app.infra.reporte_excel import _con_dia

    camion = libro["Juan Pérez"]
    assert camion["A1"].value == "El camión de Juan Pérez, día por día"
    assert camion["D4"].value == _con_dia(hoy)
    assert [c.value for c in camion[5]][3:9] == DEL_DIA
    assert _con(camion, "ATUN-140")[2:9] == ["Atún en agua 140 g", 0, 240, 180, 0, 0, 60]
    corte = _con(ventas, "Faltaron $50.00")
    assert corte[3:6] == [2250, 2200, -50]

    inventario = libro["Inventario"]
    # La bodega primero, con sus 760 del atún; el camión, con los 60 que le quedan.
    titulos = [f[0] for f in inventario.iter_rows(values_only=True) if f and f[0]]
    assert titulos.index("Bodega") < titulos.index("Camión 01 · Juan Pérez")
    atun = [f for f in inventario.iter_rows(values_only=True) if f[1] == "ATUN-140"]
    assert [f[3] for f in atun] == [760, 60]

    cargas = libro["Cargas y mermas"]
    carga = _con(cargas, "CG-TEST")
    assert carga[6:8] == [240, "Cortada"]

    cuentas = libro["Cuenta de vendedores"]
    assert _con(cuentas, "Juan Pérez")[4] == 50
    assert _con(cuentas, "Faltante de efectivo")[5] == 50

    # Con todos los encabezados en español, sin claves internas.
    textos = {c for h in libro for f in h.iter_rows(values_only=True) for c in f if c}
    assert not {"confirmada", "liquidada", "faltante_efectivo", "cargo"} & textos


async def test_lo_eliminado_no_cuenta(cliente, sesion, semilla, dia):  # noqa: F811
    await _entrar(cliente)
    await _anular(cliente, "venta", dia["venta"])
    _, libro = await _bajar(cliente)
    assert _renglon(libro["Resumen"], "Ventas de los vendedores")[1] == 0
    assert "No hubo en este periodo." in [
        f[0] for f in libro["Ventas y efectivo"].iter_rows(values_only=True)
    ]
    # El inventario es el de hoy: la venta eliminada regresó las 180 al camión.
    atun = [f for f in libro["Inventario"].iter_rows(values_only=True) if f[1] == "ATUN-140"]
    assert [Decimal(str(f[3])) for f in atun] == [760, 240]


async def test_lo_baja_la_oficina_y_sin_entrar_nadie(cliente, sesion, semilla, dia):  # noqa: F811
    # Sin entrar al panel, no hay archivo.
    r = await cliente.get("/panel/reporte/excel?periodo=hoy", follow_redirects=False)
    assert r.status_code != 200 and not r.content.startswith(b"PK")

    await _entrar(cliente, "GER01")
    _, libro = await _bajar(cliente, "semana")
    assert libro.sheetnames == HOJAS

    await _usuario(sesion, semilla, "SUP01", "supervisor")
    await _entrar(cliente, "SUP01")
    _, libro = await _bajar(cliente, "hoy")
    assert libro.sheetnames == HOJAS


# ---------------------------------------------------------------------------
# Todo es en efectivo (ADR 0002 §170)
# ---------------------------------------------------------------------------


def _textos(hoja) -> set:
    return {c for f in hoja.iter_rows(values_only=True) for c in f if c}


async def test_todo_en_efectivo_no_repite_lo_vendido(cliente, sesion, semilla, dia):  # noqa: F811
    """«Revisa que todos los pagos son solo en efectivo, incluido en los reportes
    de Excel.» El Resumen decía «Ventas de los vendedores» y abajo «De eso, en
    efectivo» con la misma cifra, como si hubiera otra forma de pago; y la hoja de
    ventas repetía «Importe» en una columna «En efectivo»."""
    await _entrar(cliente)
    _, libro = await _bajar(cliente)

    resumen = libro["Resumen"]
    vendido = _renglon(resumen, "Ventas de los vendedores")
    assert vendido[1] == 2250
    assert "todas en efectivo" in vendido[2]
    assert "De eso, en efectivo" not in _textos(resumen)

    ventas = libro["Ventas y efectivo"]
    assert "En efectivo" not in _textos(ventas)
    assert "Todas las ventas del periodo fueron en efectivo." in _textos(ventas)
    assert _con(ventas, "Juan Pérez")[2:5] == [1, 180, 2250]


async def test_una_transferencia_de_un_telefono_viejo_lo_vuelve_a_separar(
    cliente, sesion, semilla, dia  # noqa: F811
):
    """El servidor la sigue aceptando —la mercancía ya salió— (§81, enmienda):
    entonces el efectivo ya no es lo vendido, y el reporte lo dice."""
    await sesion.execute(
        text(
            "UPDATE ventas SET forma_pago = 'transferencia', pago_estado = 'por_confirmar'"
            " WHERE id = :v"
        ),
        {"v": dia["venta"]},
    )
    await sesion.commit()
    await _entrar(cliente)
    _, libro = await _bajar(cliente)

    resumen = libro["Resumen"]
    assert _renglon(resumen, "Ventas de los vendedores")[1] == 2250
    assert _renglon(resumen, "De eso, en efectivo")[1] == 0
    ventas = libro["Ventas y efectivo"]
    assert "En efectivo" in _textos(ventas)
    assert _con(ventas, "Juan Pérez")[2:6] == [1, 180, 2250, 0]


async def test_el_camion_de_cada_vendedor_dia_por_dia(cliente, sesion, semilla, dia):  # noqa: F811
    """«Quiero poder comparar por vendedor día a día, ver cómo se comporta su
    inventario» (§179). Ayer: cargó 240, vendió 180, le quedaron 60. Hoy: traía
    esos 60, cargó 120, vendió 100, mermó 10, la oficina le quitó 5 y le quedaron
    65. Cada día en sus columnas, uno al lado del otro."""
    import uuid
    from datetime import timedelta

    from app.infra.reporte_excel import _con_dia

    hoy = dia["dia"]
    ayer = hoy - timedelta(days=1)
    await sesion.execute(text("UPDATE cargas SET fecha_operativa = :d WHERE id = :c"),
                         {"d": ayer, "c": dia["carga"]})
    await sesion.execute(text("UPDATE ventas SET fecha_operativa = :d WHERE id = :v"),
                         {"d": ayer, "v": dia["venta"]})
    carga, venta, merma = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    datos = {"carga": carga, "venta": venta, "merma": merma, "hoy": hoy,
             "b": semilla["bodega"], "c": semilla["camion"], "v": semilla["vendedor"],
             "r": semilla["ruta"], "p": dia["producto"], "d": dia["dispositivo"],
             "cli": dia["cliente"], "l": semilla["lista_precios"]}
    for sql in (
        # A $12.50 la pieza, para el camión en dinero (§185).
        "INSERT INTO precios (lista_id, producto_id, unidad_codigo, precio) "
        "VALUES (:l, :p, 'PZA', 12.50)",
        "INSERT INTO cargas (id, folio, almacen_origen_id, almacen_destino_id, vendedor_id, "
        "  ruta_id, fecha_operativa, estado, confirmada_en) "
        "VALUES (:carga, 'CG-HOY', :b, :c, :v, :r, :hoy, 'confirmada', now())",
        "INSERT INTO carga_detalle (carga_id, producto_id, cantidad) VALUES (:carga, :p, 120)",
        "INSERT INTO movimientos_inventario (tipo, almacen_origen_id, almacen_destino_id, "
        "  producto_id, cantidad, documento_tipo, documento_id) "
        "VALUES ('carga', :b, :c, :p, 120, 'carga', :carga)",
        "INSERT INTO ventas (id, dispositivo_id, folio_consecutivo, folio_local, cliente_id, "
        "  vendedor_id, almacen_id, carga_id, tipo, subtotal, total, fecha_dispositivo, "
        "  fecha_operativa) "
        "VALUES (:venta, :d, 2, 'VEND01-000002', :cli, :v, :c, :carga, 'contado', 1250, "
        "        1250, now(), :hoy)",
        "INSERT INTO venta_partidas (id, venta_id, linea, producto_id, unidad_codigo, "
        "  factor_unidad, cantidad, cantidad_base, precio_unitario, importe) "
        "VALUES (gen_random_uuid(), :venta, 1, :p, 'PZA', 1, 100, 100, 12.5, 1250)",
        "INSERT INTO movimientos_inventario (tipo, almacen_origen_id, producto_id, cantidad, "
        "  documento_tipo, documento_id) VALUES ('venta', :c, :p, 100, 'venta', :venta)",
        "INSERT INTO mermas (id, dispositivo_id, folio_consecutivo, folio_local, tipo, "
        "  almacen_id, vendedor_id, motivo_codigo, fecha_dispositivo, fecha_operativa) "
        "VALUES (:merma, :d, 1, 'M-HOY', 'merma', :c, :v, 'ROTO', now(), :hoy)",
        "INSERT INTO merma_detalle (id, merma_id, producto_id, cantidad_base) "
        "VALUES (gen_random_uuid(), :merma, :p, 10)",
        "INSERT INTO movimientos_inventario (tipo, almacen_origen_id, producto_id, cantidad, "
        "  documento_tipo, documento_id) VALUES ('merma', :c, :p, 10, 'merma', :merma)",
        # Un ajuste de la oficina: 5 menos.
        "INSERT INTO movimientos_inventario (tipo, almacen_origen_id, producto_id, cantidad, "
        "  documento_tipo, documento_id) "
        "VALUES ('ajuste', :c, :p, 5, 'ajuste_camion', gen_random_uuid())",
        "UPDATE existencias SET cantidad = cantidad + 120 - 100 - 10 - 5 "
        " WHERE almacen_id = :c AND producto_id = :p",
    ):
        await sesion.execute(text(sql), datos)
    await sesion.commit()
    await _entrar(cliente)

    _, libro = await _bajar(cliente, f"rango&desde={ayer}&hasta={hoy}")

    camion = libro["Juan Pérez"]
    assert camion["D4"].value == _con_dia(ayer)
    assert camion["J4"].value == _con_dia(hoy)
    assert [c.value for c in camion[5]][3:15] == DEL_DIA * 2
    atun = _con(camion, "ATUN-140")
    assert atun[3:9] == [0, 240, 180, 0, 0, 60]
    assert atun[9:15] == [60, 120, 100, 10, -5, 65]
    assert _con(camion, "Total")[9:15] == [60, 120, 100, 10, -5, 65]

    # «Que el Excel muestre el valor del carro del vendedor cada día» (§185).
    dinero = _con(camion, "En dinero")
    assert dinero[3:9] == [0, 3000, 2250, 0, 0, 750]
    assert dinero[9:15] == [750, 1500, 1250, 125, -62.5, 812.5]


async def test_dia_por_dia_el_dinero_y_la_mercancia_con_el_mayoreo_aparte(
    cliente, sesion, semilla, dia  # noqa: F811
):
    """«Quiero en el Excel ser capaz de distinguir que todo va bien día a día;
    separa la venta a mayoreo de la de los vendedores» (§184), y «cuánto hubo de
    dinero en todo el producto cada fecha, por camión y por almacén, y la suma»
    (§183).

    Hasta ayer la bodega tenía 1,000 atunes, a $12.50: $12,500. Hoy entraron 100
    ($1,250), se cargaron 240 al camión, el vendedor vendió 200 ($2,500) y entregó
    $2,400, y la oficina vendió una caja de 24 a mayoreo en $270 ($300 de lista)."""
    import uuid
    from datetime import timedelta

    from app.infra import mayoreo
    from app.infra.reporte_excel import _con_dia
    from tests.test_corte_sin_bajar_la_carga import _vende_el_sabado

    hoy = dia["dia"]
    ayer, anteayer = hoy - timedelta(days=1), hoy - timedelta(days=2)
    await sesion.execute(
        text("INSERT INTO precios (lista_id, producto_id, unidad_codigo, precio) "
             "VALUES (:l, :p, 'PZA', 12.50)"),
        {"l": semilla["lista_precios"], "p": dia["producto"]},
    )
    # La compra que llegó a la bodega.
    await sesion.execute(
        text("INSERT INTO movimientos_inventario (tipo, almacen_destino_id, producto_id, "
             "  cantidad, documento_tipo, documento_id) "
             "VALUES ('compra', :b, :p, 100, 'entrada', :doc)"),
        {"b": semilla["bodega"], "p": dia["producto"], "doc": uuid.uuid4()},
    )
    await sesion.execute(
        text("UPDATE existencias SET cantidad = cantidad + 100 "
             " WHERE almacen_id = :b AND producto_id = :p"),
        {"b": semilla["bodega"], "p": dia["producto"]},
    )
    await sesion.commit()
    await _vende_el_sabado(sesion, semilla, dia, carga=dia["carga"], piezas=20)
    corte = await _corte(sesion, dia, semilla, efectivo="2400.00")
    await _cerrar(cliente, corte, await _cab(cliente, "GER01"))
    nuevo = await mayoreo.alta_de_cliente(sesion, nombre="Abarrotes El Mayorista",
                                          quien=semilla["admin"])
    await mayoreo.vender(
        sesion, cliente_id=nuevo["id"], almacen_id=semilla["bodega"],
        renglones=[mayoreo.RenglonPedido(producto_id=dia["producto"], unidad_codigo="CAJA",
                                         bultos=Decimal(1), precio=Decimal(270))],
        forma_pago="efectivo", referencia_pago=None, nota=None, quien=semilla["admin"],
    )
    await sesion.commit()
    await _entrar(cliente)

    _, libro = await _bajar(cliente, f"rango&desde={anteayer}&hasta={hoy + timedelta(days=1)}")
    hoja = libro["Día por día"]

    # El dinero: lo de los vendedores y su corte, y el mayoreo aparte.
    assert _con(hoja, "Vendieron los vendedores")[:6] == [
        "Día", "Vendieron los vendedores", "Entregaron en sus cortes", "Diferencia",
        "Venta de mayoreo", "Total vendido"]
    filas = [list(f) for f in hoja.iter_rows(values_only=True) if f and f[0] == _con_dia(hoy)]
    dinero, mercancia, almacenes = filas
    assert dinero[1:6] == [2500, 2400, -100, 270, 2770]
    assert _renglon(hoja, _con_dia(ayer))[1:6] == [0, 0, 0, 0, 0]

    # La mercancía: se suma de izquierda a derecha.
    assert _con(hoja, "Valía al empezar")[:10] == [
        "Día", "Valía al empezar", "Entró a la bodega", "Vendieron los vendedores",
        "Venta de mayoreo", "Mermas", "Devoluciones de clientes", "Ajustes y conteos",
        "Eliminado en Operaciones", "Valía al cerrar"]
    assert mercancia[1:10] == [12500, 1250, -2500, -300, 0, 0, 0, 0, 10950]
    totales = [list(f) for f in hoja.iter_rows(values_only=True) if f and f[0] == "Total"]
    assert totales[0][1:6] == [2500, 2400, -100, 270, 2770]
    assert totales[1][1:10] == [12500, 1250, -2500, -300, 0, 0, 0, 0, 10950]

    # Cada almacén y la suma; la carga solo cambió de lugar.
    assert _con(hoja, "Bodega")[:4] == ["Día", "Bodega", "Camión 01 · Juan Pérez", "Total"]
    assert almacenes[1:4] == [10450, 500, 10950]
    del_anteayer = [list(f) for f in hoja.iter_rows(values_only=True)
                    if f and f[0] == _con_dia(anteayer)]
    assert del_anteayer[2][1:4] == [12500, 0, 12500]
    # Hasta hoy: de mañana no se sabe nada.
    assert not any(f and f[0] == _con_dia(hoy + timedelta(days=1))
                   for f in hoja.iter_rows(values_only=True))
